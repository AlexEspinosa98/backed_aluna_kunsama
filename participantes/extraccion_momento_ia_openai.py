"""Lee un Momento ya diligenciado fuera de la web (PDF o Word) y lo transcribe con UNA sola
llamada a OpenAI — mismo mecanismo que instrumentos.extraccion_ia_openai (ver ese módulo para el
razonamiento completo de texto-vs-visión), portado acá porque este contenido vive en el modelo
clásico `jornadas.Momento`/`jornadas.Pregunta`, no en el módulo `instrumentos`.

El esquema que se le manda a la IA lleva, por pregunta, su id real, su tipo, sus opciones y —en
las de tabla— sus filas y columnas con id propio. Transcribe los seis tipos de `Pregunta`,
incluidas las FILAS AGREGADAS: si una pregunta tiene `filas_adicionales` (las filas extra de una
matriz, y todas las filas de una lista), la IA puede transcribir filas que no están en el esquema
agrupándolas con un `fila_temporal` que ella misma inventa — el mismo mecanismo que usa el envío
normal desde la web (ver participantes.views).

El resultado queda guardado en `ExtraccionMomento.resultado` (JSON) SIN tocar `Respuesta` — la
escritura real pasa por `aprobar_extraccion_momento`, disparada a mano por un admin desde la
vista (ver participantes/admin_views.py)."""
import base64
import json
import os
import threading

from django.db import transaction
from django.utils import timezone

from jornadas import emparejamiento

DEFAULT_MODEL = os.environ.get('OPENAI_MODEL', 'gpt-4o')
REASONING_EFFORT = os.environ.get('OPENAI_REASONING_EFFORT', 'medium')
GENERATION_TIMEOUT_SECONDS = 300
MAX_OUTPUT_TOKENS = 8000
MODELO_USADO_LABEL = 'Generado con IA'
MAX_PAGINAS_IMAGEN = 20
UMBRAL_CARACTERES_POR_PAGINA = 40

SYSTEM_PROMPT_EXTRACCION = (
    "Eres un transcriptor de formularios. Se te entrega, en JSON, el ESQUEMA completo de un "
    "momento (lista de preguntas, cada una con su id real, tipo y — si aplica — sus opciones con "
    "su id real), y el contenido de un documento (PDF o Word) que alguien ya diligenció a mano o "
    "en computador fuera de este sistema. Tu única tarea es leer ese documento y transcribir, "
    "para cada pregunta que SÍ tenga contenido en el documento, una entrada que apunte al id "
    "real de esa pregunta.\n\n"

    "=== REGLAS ESTRICTAS ===\n"
    "1. USA SOLO los ids que te dimos en el esquema — nunca inventes uno.\n"
    "2. Si una pregunta no tiene contenido en el documento (en blanco, ilegible, no la "
    "encuentras), NO incluyas una entrada para ella — no inventes ni dejes texto vacío como "
    "relleno.\n"
    "3. Pregunta tipo 'abierta' o 'audio': una entrada con texto_libre = lo transcrito "
    "literalmente (puedes corregir errores obvios de tipeo, pero no resumas ni parafrasees). "
    "Las dos se tratan igual acá: en un documento diligenciado una pregunta 'audio' ya viene "
    "escrita, no grabada.\n"
    "4. Pregunta tipo 'unica'/'multiple': una entrada con opcion_ids = [id de la(s) opción(es) "
    "marcada(s)], usando el id de la opción cuyo texto mejor coincida con lo marcado/escrito — "
    "nunca opciones que no estén en el esquema de esa pregunta.\n"
    "5. Pregunta tipo 'matriz': UNA entrada POR CADA celda (fila×columna) que tenga contenido, "
    "con fila_id y columna_id = los ids correspondientes de ESA pregunta y texto_libre = lo "
    "escrito en esa celda — nunca mezcles filas/columnas de una matriz con otra.\n"
    "6. FILAS AGREGADAS. Algunas preguntas traen 'filas_adicionales': true en el esquema. Ahí el "
    "documento puede tener filas que NO están en la lista 'filas' — se las agregó a mano quien "
    "diligenció (una fila extra al final de la tabla, un renglón escrito en el margen, una hoja "
    "anexa con más registros). Transcríbelas así:\n"
    "   - una entrada POR CADA celda de esa fila nueva, con columna_id = la columna que "
    "corresponde y texto_libre = lo escrito;\n"
    "   - en vez de fila_id, pon 'fila_temporal': un número que TÚ inventas solo para agrupar "
    "las celdas de una misma fila nueva (1 para la primera fila agregada de esa pregunta, 2 para "
    "la segunda...). NO es el id de nada: no lo busques en el esquema, no lo reutilices entre "
    "preguntas distintas;\n"
    "   - nunca pongas fila_id y fila_temporal en la misma entrada: o la fila ya existía en el "
    "esquema (fila_id) o la agregó quien diligenció (fila_temporal).\n"
    "   Una pregunta tipo 'lista' funciona SOLO así: no tiene filas predefinidas, todas sus "
    "filas van con fila_temporal.\n"
    "   Si la pregunta NO trae 'filas_adicionales': true, no uses fila_temporal ahí — cualquier "
    "fila que no esté en el esquema simplemente no se transcribe.\n"
    "7. Nunca mezcles: no le pongas opcion_ids a una pregunta abierta/audio, matriz o lista, no le "
    "pongas fila_id/columna_id a una pregunta que no sea matriz/lista, ni texto_libre a una de "
    "opción.\n"
    "8. RESPONSABLE. Además de las respuestas, busca quién diligenció el documento — casi siempre "
    "está en las primeras líneas o en un encabezado, con etiquetas como 'Responsable', "
    "'Diligenciado por', 'Nombre', 'Elaborado por' o una firma al pie. Devuélvelo en el "
    "bloque 'responsable', copiando el texto TAL CUAL aparece (no lo normalices, no lo "
    "completes, no deduzcas un correo que no esté escrito). Si un dato no está, déjalo en null; "
    "si no encuentras ningún responsable, deja todo el bloque en null. NUNCA inventes un nombre "
    "ni uses el de una persona mencionada dentro de una respuesta — solo quien firma o declara "
    "haber diligenciado el documento.\n\n"

    "=== FORMATO DE SALIDA (obligatorio) ===\n"
    "Responde ÚNICAMENTE con un objeto JSON válido, sin explicación antes ni después, sin fences "
    "de markdown, con esta forma exacta:\n"
    "{\n"
    '  "responsable": {"nombre": "<texto o null>", "correo": "<texto o null>", '
    '"cargo": "<texto o null>", "dependencia": "<texto o null>"},\n'
    '  "respuestas": [\n'
    '    {"pregunta": <id>, "texto_libre": "<texto o \'\'>", "opcion_ids": [<ids>], '
    '"fila_id": <id o null>, "fila_temporal": <número o null>, "columna_id": <id o null>}\n'
    "  ]\n"
    "}\n"
)


def _extraer_texto_docx(archivo):
    from docx import Document

    documento = Document(archivo)
    partes = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            if any(celdas):
                partes.append(' | '.join(celdas))
    return '\n'.join(partes)


def _extraer_texto_o_imagenes_pdf(archivo):
    import pdfplumber

    archivo.seek(0)
    paginas_texto = []
    with pdfplumber.open(archivo) as pdf:
        n_paginas = len(pdf.pages)
        for pagina in pdf.pages:
            paginas_texto.append(pagina.extract_text() or '')

    texto = '\n'.join(paginas_texto).strip()
    if n_paginas and len(texto) / n_paginas >= UMBRAL_CARACTERES_POR_PAGINA:
        return texto, None

    import fitz  # PyMuPDF

    archivo.seek(0)
    imagenes = []
    documento = fitz.open(stream=archivo.read(), filetype='pdf')
    for pagina in documento[:MAX_PAGINAS_IMAGEN]:
        pixmap = pagina.get_pixmap(dpi=150)
        imagenes.append(base64.b64encode(pixmap.tobytes('png')).decode('ascii'))
    documento.close()
    return None, imagenes


def _leer_documento(extraccion):
    nombre = extraccion.nombre_archivo_original or extraccion.archivo.name
    extension = nombre.rsplit('.', 1)[-1].lower() if '.' in nombre else ''

    extraccion.archivo.open('rb')
    try:
        if extension == 'docx':
            return _extraer_texto_docx(extraccion.archivo), None
        if extension == 'pdf':
            return _extraer_texto_o_imagenes_pdf(extraccion.archivo)
        raise ValueError(f'Formato de archivo no soportado: .{extension} (solo .pdf o .docx).')
    finally:
        extraccion.archivo.close()


def _construir_payload_esquema(momento):
    from jornadas.models import Pregunta

    preguntas = []
    for pregunta in momento.preguntas.filter(activa=True).order_by('orden'):
        # Las columnas van tanto en matriz como en lista (una lista ES columnas fijas + filas
        # dinámicas). Las filas predefinidas solo existen en matriz. `filas_adicionales` es lo que
        # le dice a la IA si puede transcribir filas que no están en el esquema, usando
        # fila_temporal — sin este dato en el payload la regla 6 del prompt no tendría cómo
        # aplicarse pregunta por pregunta.
        es_tabla = pregunta.tipo in (Pregunta.TIPO_MATRIZ, Pregunta.TIPO_LISTA)
        preguntas.append({
            'id': pregunta.id,
            'texto': pregunta.texto,
            'tipo': pregunta.tipo,
            'filas_adicionales': pregunta.acepta_filas_dinamicas,
            'opciones': [{'id': o.id, 'texto': o.texto} for o in pregunta.opciones.all()],
            'filas': (
                [{'id': f.id, 'texto': f.texto} for f in pregunta.filas.all()]
                if pregunta.tipo == Pregunta.TIPO_MATRIZ else []
            ),
            'columnas': (
                [{'id': c.id, 'texto': c.texto} for c in pregunta.columnas.all()]
                if es_tabla else []
            ),
        })
    return {'momento': momento.titulo, 'preguntas': preguntas}


def _llamar_openai_extraccion(esquema, texto_documento=None, imagenes_base64=None):
    api_key = os.environ.get('OPENAI_API_KEY')
    if not api_key:
        return None, 'OPENAI_API_KEY no está configurada en el entorno del servidor (.env).'

    encabezado = 'ESQUEMA DEL MOMENTO (JSON):\n' + json.dumps(esquema, ensure_ascii=False, indent=2)

    if imagenes_base64:
        contenido_usuario = [
            {'type': 'text', 'text': encabezado + '\n\nDOCUMENTO DILIGENCIADO (páginas adjuntas como imagen):'},
        ]
        for imagen_b64 in imagenes_base64:
            contenido_usuario.append({
                'type': 'image_url',
                'image_url': {'url': f'data:image/png;base64,{imagen_b64}', 'detail': 'high'},
            })
    else:
        contenido_usuario = (
            encabezado + '\n\nDOCUMENTO DILIGENCIADO (texto extraído):\n' + (texto_documento or '')
        )

    resultado = {}

    def _run():
        try:
            from openai import OpenAI
            client = OpenAI(api_key=api_key)
            kwargs = dict(
                model=DEFAULT_MODEL,
                messages=[
                    {'role': 'system', 'content': SYSTEM_PROMPT_EXTRACCION},
                    {'role': 'user', 'content': contenido_usuario},
                ],
                max_completion_tokens=MAX_OUTPUT_TOKENS,
                response_format={'type': 'json_object'},
            )
            if REASONING_EFFORT:
                kwargs['reasoning_effort'] = REASONING_EFFORT
            else:
                kwargs['temperature'] = 0.2
            respuesta = client.chat.completions.create(**kwargs)
            resultado['texto'] = respuesta.choices[0].message.content.strip()
        except Exception as exc:  # noqa: BLE001
            resultado['error'] = str(exc)

    hilo = threading.Thread(target=_run, daemon=True)
    hilo.start()
    hilo.join(timeout=GENERATION_TIMEOUT_SECONDS)

    if hilo.is_alive():
        return None, f'Tiempo de espera agotado ({GENERATION_TIMEOUT_SECONDS}s) esperando a OpenAI.'
    if resultado.get('error'):
        return None, resultado['error']
    texto = resultado.get('texto')
    if not texto:
        return None, 'OpenAI no devolvió contenido.'
    try:
        return json.loads(texto), None
    except json.JSONDecodeError as exc:
        return None, f'OpenAI devolvió JSON inválido: {exc}'


def _limpiar_y_validar(resultado_crudo, momento):
    """Descarta cualquier entrada que no pase la misma validación por tipo que ya usa el envío
    normal (participantes.views._validar_entrada) — a diferencia de ese endpoint, NO aborta todo
    si una entrada falla, solo la omite (queda en preguntas_omitidas). No escribe Respuesta acá —
    solo deja `resultado` listo para que aprobar_extraccion_momento lo escriba después."""
    from jornadas.models import Pregunta

    from participantes.views import _validar_entrada

    preguntas_validas = {p.id: p for p in momento.preguntas.filter(activa=True)}
    limpio = []
    omitidas = []
    for crudo in resultado_crudo.get('respuestas') or []:
        pregunta = preguntas_validas.get(crudo.get('pregunta'))
        if pregunta is None:
            omitidas.append(crudo.get('pregunta'))
            continue
        texto_libre = crudo.get('texto_libre', '') or ''
        opcion_ids = crudo.get('opcion_ids') or []
        opciones = list(pregunta.opciones.filter(id__in=opcion_ids))

        fila = columna = None
        fila_temporal = None
        if pregunta.tipo in (Pregunta.TIPO_MATRIZ, Pregunta.TIPO_LISTA):
            columna = pregunta.columnas.filter(id=crudo.get('columna_id')).first()
            if pregunta.tipo == Pregunta.TIPO_MATRIZ:
                fila = pregunta.filas.filter(id=crudo.get('fila_id')).first()
            if pregunta.acepta_filas_dinamicas and crudo.get('fila_temporal') is not None:
                # `fila_temporal` lo inventa la IA para agrupar las celdas de una fila que no
                # estaba en el esquema (ver regla 6 del prompt). No se resuelve contra la base
                # —no apunta a nada todavía—, solo tiene que ser un entero; cualquier otra cosa
                # es una alucinación y la entrada se omite como cualquier otra inconsistencia.
                try:
                    fila_temporal = int(crudo['fila_temporal'])
                except (TypeError, ValueError):
                    omitidas.append(pregunta.id)
                    continue

        try:
            _validar_entrada(pregunta, texto_libre, opciones, fila, columna, fila_temporal)
        except Exception:  # noqa: BLE001 — ValidationError u otra inconsistencia de la IA
            omitidas.append(pregunta.id)
            continue
        limpio.append({
            'pregunta': pregunta.id, 'texto_libre': texto_libre, 'opcion_ids': [o.id for o in opciones],
            'fila_id': fila.id if fila else None, 'fila_temporal': fila_temporal,
            'columna_id': columna.id if columna else None,
        })

    return {'respuestas': limpio}, omitidas


def _limpiar_responsable(crudo):
    """Normaliza el bloque `responsable` que devolvió la IA a un dict de strings. La IA puede
    mandar null, omitirlo, o mandar cualquier cosa en los campos — nada de eso puede tumbar una
    extracción que por lo demás salió bien."""
    if not isinstance(crudo, dict):
        return {}
    limpio = {}
    for clave in ('nombre', 'correo', 'cargo', 'dependencia'):
        valor = crudo.get(clave)
        if isinstance(valor, str) and valor.strip():
            limpio[clave] = valor.strip()
    return limpio


def emparejar_responsable_momento(momento, responsable):
    """Busca a qué Participante de ESTA jornada corresponde el responsable que leyó la IA.

    El conjunto de candidatos es la jornada completa, no el momento: un participante existe a
    nivel de jornada y puede no haber respondido todavía nada de este momento — que es
    justamente el caso de quien llenó el documento en papel."""
    from .models import Participante

    if not responsable:
        return None, emparejamiento.ESTADO_SIN_DATO

    candidatos = [
        (p, f'{p.nombre} {p.apellido}', p.correo_institucional)
        for p in Participante.objects.filter(jornada_id=momento.jornada_id)
    ]
    return emparejamiento.emparejar(
        candidatos, nombre=responsable.get('nombre'), correo=responsable.get('correo'),
    )


def procesar_extraccion_momento(extraccion_id):
    """Genera el `resultado` de una ExtraccionMomento ya creada (estado 'pendiente'). Corre en un
    hilo de background — mismo patrón que instrumentos.extraccion_ia_openai. NO escribe
    Respuesta — deja el resultado listo para revisión/aprobación manual."""
    from django.db import close_old_connections

    close_old_connections()
    from .models import ExtraccionMomento

    extraccion = None
    try:
        extraccion = ExtraccionMomento.objects.select_related('momento').get(pk=extraccion_id)
        extraccion.estado = ExtraccionMomento.ESTADO_PROCESANDO
        extraccion.save(update_fields=['estado'])

        texto, imagenes = _leer_documento(extraccion)
        esquema = _construir_payload_esquema(extraccion.momento)
        resultado, error = _llamar_openai_extraccion(esquema, texto_documento=texto, imagenes_base64=imagenes)

        if resultado is None:
            extraccion.estado = ExtraccionMomento.ESTADO_ERROR
            extraccion.error_mensaje = error or 'Error desconocido generando la extracción.'
            extraccion.save(update_fields=['estado', 'error_mensaje'])
            return

        limpio, omitidas = _limpiar_y_validar(resultado, extraccion.momento)

        # El responsable solo se busca si no lo dijeron al subir: un participante indicado a mano
        # es una decisión humana y no se pisa con lo que haya leído la IA.
        campos_responsable = []
        if extraccion.participante_id is None:
            detectado = _limpiar_responsable(resultado.get('responsable'))
            participante, estado = emparejar_responsable_momento(extraccion.momento, detectado)
            extraccion.responsable_detectado = detectado
            extraccion.responsable_estado = estado
            if participante is not None:
                extraccion.participante = participante
            campos_responsable = ['responsable_detectado', 'responsable_estado', 'participante']

        extraccion.resultado = limpio
        extraccion.preguntas_omitidas = omitidas
        extraccion.estado = ExtraccionMomento.ESTADO_COMPLETO
        extraccion.error_mensaje = ''
        extraccion.modelo_usado = MODELO_USADO_LABEL
        extraccion.completado_en = timezone.now()
        extraccion.save(update_fields=[
            'resultado', 'preguntas_omitidas', 'estado', 'error_mensaje', 'modelo_usado',
            'completado_en', *campos_responsable,
        ])
    except Exception as exc:  # noqa: BLE001
        if extraccion is not None:
            extraccion.estado = ExtraccionMomento.ESTADO_ERROR
            extraccion.error_mensaje = str(exc)
            extraccion.save(update_fields=['estado', 'error_mensaje'])
    finally:
        close_old_connections()


def _escribir_filas_agregadas(pregunta, items, participante):
    """Escribe las celdas que la IA transcribió como filas AGREGADAS (`fila_temporal`): las filas
    extra de una matriz con filas_adicionales, y todas las filas de una lista.

    Reemplaza las filas dinámicas que ese participante ya tuviera **en esta pregunta**, igual que
    el envío normal (participantes.views._guardar_filas_dinamicas) — el documento aprobado es la
    versión buena de esa tabla, no un anexo a lo anterior. A diferencia del envío normal, acá NO
    se borran las filas de preguntas que la IA no mencionó: aprobar una extracción escribe lo que
    el documento traía, y un documento que no habla de una pregunta no es una instrucción de
    borrarla."""
    from .models import FilaListaRespuesta, Respuesta

    FilaListaRespuesta.objects.filter(pregunta=pregunta, participante=participante).delete()

    filas_por_temporal = {
        ft: FilaListaRespuesta.objects.create(
            pregunta=pregunta, participante=participante, orden=i,
        )
        for i, ft in enumerate(sorted({item['fila_temporal'] for item in items}), start=1)
    }

    guardadas = []
    for item in items:
        guardadas.append(Respuesta.objects.create(
            pregunta=pregunta,
            participante=participante,
            fila_lista=filas_por_temporal[item['fila_temporal']],
            columna_id=item.get('columna_id'),
            texto_libre=item.get('texto_libre', ''),
            registrado_por=participante,
        ))
    return guardadas


@transaction.atomic
def aprobar_extraccion_momento(extraccion, aprobado_por):
    """Escribe `extraccion.resultado` como Respuesta reales del participante — mismo lookup/save
    que RespuestasMomentoView.post() (participante individual, registrado_por=el mismo
    participante). Solo llamable sobre una extracción en estado completo y no aprobada aún (ver
    la vista para esas guardas).

    Atómico por lo mismo que el envío normal: las filas agregadas se borran y se recrean, así que
    una falla a mitad dejaría la tabla del participante incompleta."""
    from django.utils import timezone as tz

    from rest_framework.exceptions import ValidationError

    from jornadas.models import Pregunta

    from .models import Respuesta

    if extraccion.participante_id is None:
        # Desde HU-55 el responsable puede venir sin resolver (la IA no lo encontró, o encontró
        # dos personas con ese nombre). La transcripción sigue intacta en `resultado`; lo único
        # que falta es a nombre de quién se escribe, y eso lo decide una persona.
        raise ValidationError({'participante': (
            'Esta extracción todavía no tiene responsable asignado. Asigna uno antes de aprobarla '
            f'(responsable detectado: {extraccion.responsable_detectado or "ninguno"}).'
        )})

    preguntas = {p.id: p for p in Pregunta.objects.filter(momento=extraccion.momento)}
    guardadas = []
    # Las celdas de fila agregada se juntan por pregunta antes de escribir: una FilaListaRespuesta
    # agrupa varias celdas, así que no se puede ir celda por celda como con las de fila fija.
    agregadas_por_pregunta = {}

    for item in extraccion.resultado.get('respuestas') or []:
        pregunta = preguntas.get(item['pregunta'])
        if pregunta is None:
            continue
        if item.get('fila_temporal') is not None and pregunta.acepta_filas_dinamicas:
            agregadas_por_pregunta.setdefault(pregunta, []).append(item)
            continue
        respuesta, _ = Respuesta.objects.update_or_create(
            pregunta=pregunta, participante=extraccion.participante,
            fila_id=item.get('fila_id'), columna_id=item.get('columna_id'),
            defaults={'texto_libre': item.get('texto_libre', ''), 'registrado_por': extraccion.participante},
        )
        respuesta.opciones.set(item.get('opcion_ids') or [])
        guardadas.append(respuesta)

    for pregunta, items in agregadas_por_pregunta.items():
        guardadas.extend(_escribir_filas_agregadas(pregunta, items, extraccion.participante))

    extraccion.aprobado_en = tz.now()
    extraccion.aprobado_por = aprobado_por
    extraccion.save(update_fields=['aprobado_en', 'aprobado_por'])
    return guardadas
