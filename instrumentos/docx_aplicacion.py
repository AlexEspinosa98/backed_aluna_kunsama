"""Genera el .docx de una AplicacionInstrumento ya diligenciada — descarga habilitada tanto para
el encargado (AplicacionInstrumentoAdminViewSet.descargar) como para el propio preregistrado
(InstrumentoRespuestasView.get), siempre a partir de datos ya guardados en BD, nunca generado por
IA (mismo principio que analitica/pdf_presentacion.py: los documentos de salida reflejan lo que ya
está calculado/almacenado)."""
import io
from collections import defaultdict

from django.http import HttpResponse
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .models import PreguntaInstrumento


def _agrupar_respuestas(aplicacion):
    por_pregunta = defaultdict(list)
    for respuesta in aplicacion.respuestas.select_related('fila', 'columna').prefetch_related('opciones'):
        por_pregunta[respuesta.pregunta_id].append(respuesta)
    return por_pregunta


def _escribir_pregunta(document, pregunta, respuestas):
    parrafo = document.add_paragraph()
    run = parrafo.add_run(pregunta.texto)
    run.bold = True

    if pregunta.tipo == PreguntaInstrumento.TIPO_MATRIZ:
        filas = list(pregunta.filas.all())
        columnas = list(pregunta.columnas.all())
        if not filas or not columnas:
            document.add_paragraph('(matriz sin filas/columnas definidas)')
            return

        celdas = {(r.fila_id, r.columna_id): r.texto_libre for r in respuestas}
        tabla = document.add_table(rows=len(filas) + 1, cols=len(columnas) + 1)
        tabla.style = 'Table Grid'
        tabla.cell(0, 0).text = ''
        for c_idx, columna in enumerate(columnas, start=1):
            tabla.cell(0, c_idx).text = columna.texto
        for f_idx, fila in enumerate(filas, start=1):
            tabla.cell(f_idx, 0).text = fila.texto
            for c_idx, columna in enumerate(columnas, start=1):
                tabla.cell(f_idx, c_idx).text = celdas.get((fila.id, columna.id), '') or ''
        return

    respuesta = respuestas[0] if respuestas else None
    if pregunta.tipo == PreguntaInstrumento.TIPO_ABIERTA:
        texto = respuesta.texto_libre.strip() if respuesta else ''
        document.add_paragraph(texto or '(sin respuesta)')
    else:
        opciones = ', '.join(o.texto for o in respuesta.opciones.all()) if respuesta else ''
        document.add_paragraph(opciones or '(sin respuesta)')


def generar_docx_aplicacion(aplicacion):
    preregistro = aplicacion.preregistro
    instrumento = preregistro.instrumento
    usuario = preregistro.usuario

    document = Document()

    titulo = document.add_heading(instrumento.nombre, level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = document.add_paragraph()
    nombre_completo = f'{usuario.first_name} {usuario.last_name}'.strip() or usuario.username
    meta.add_run(f'Diligenciado por: {nombre_completo} ({usuario.email or usuario.username})\n').bold = True
    fecha = timezone.localtime(aplicacion.enviado_en) if aplicacion.enviado_en else None
    meta.add_run(f'Enviado: {fecha.strftime("%Y-%m-%d %H:%M") if fecha else "—"}\n')
    meta.add_run(f'Estado: {aplicacion.get_estado_display() if aplicacion.enviado_en else "Sin enviar"}\n')
    if aplicacion.comentario_revision:
        meta.add_run(f'Comentario de revisión: {aplicacion.comentario_revision}\n')

    respuestas_por_pregunta = _agrupar_respuestas(aplicacion)

    for seccion in instrumento.secciones.filter(activa=True).order_by('orden'):
        document.add_heading(seccion.titulo, level=1)
        if seccion.tipo == seccion.TIPO_CONTENIDO:
            for parrafo in seccion.contenido.split('\n'):
                if parrafo.strip():
                    document.add_paragraph(parrafo.strip())
            continue

        for pregunta in seccion.preguntas.filter(activa=True).order_by('orden'):
            _escribir_pregunta(document, pregunta, respuestas_por_pregunta.get(pregunta.id, []))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def respuesta_docx_http(aplicacion):
    contenido = generar_docx_aplicacion(aplicacion)
    response = HttpResponse(
        contenido,
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    usuario = aplicacion.preregistro.usuario
    instrumento_slug = aplicacion.preregistro.instrumento.slug
    fecha = timezone.localtime(timezone.now()).strftime('%Y%m%d')
    response['Content-Disposition'] = (
        f'attachment; filename="aplicacion-{instrumento_slug}-{usuario.username}-{fecha}.docx"'
    )
    return response
