"""Lee un instrumento ya diligenciado fuera de la web (PDF o Word) y lo transcribe con UNA sola
llamada a OpenAI al MISMO formato que ya usa el envío normal desde la web
(RespuestaInstrumentoEnvioSerializer, ver views_participante.py) — la IA no inventa una
estructura propia, solo "llena el mismo formulario" que llenaría la persona, incluyendo los ids
reales de cada pregunta/opción/fila/columna (se le mandan en el payload, nunca los adivina).

Dos vías de lectura del documento, elegidas automáticamente:
- Word (.docx) o PDF con texto seleccionable: se extrae el texto plano (python-docx / pdfplumber)
  y se manda como texto al modelo — más barato y más preciso.
- PDF escaneado o diligenciado a mano (el texto extraído sale vacío o casi vacío): se renderizan
  las páginas como imágenes (PyMuPDF) y se mandan con capacidad de visión (gpt-4o) — mismo modelo,
  solo cambia qué contenido lleva el mensaje.

El resultado nunca se acepta solo: cae siempre en AplicacionInstrumento.estado='pendiente' con
generado_por_ia=True para revisión humana (ver ExtraccionInstrumento en models.py)."""
import base64
import json
import os
import threading

from django.utils import timezone

from jornadas import emparejamiento

DEFAULT_MODEL = os.environ.get('OPENAI_MODEL', 'gpt-4o')
REASONING_EFFORT = os.environ.get('OPENAI_REASONING_EFFORT', 'medium')
GENERATION_TIMEOUT_SECONDS = 300
MAX_OUTPUT_TOKENS = 8000
# Nunca se expone el nombre real del modelo de un proveedor externo en la respuesta de la API.
MODELO_USADO_LABEL = 'Generado con IA'
# Tope de páginas enviadas como imagen — un PDF escaneado de 90 preguntas rara vez pasa de esto, y
# limita el costo/tiempo de una sola llamada con visión.
MAX_PAGINAS_IMAGEN = 20
# Si el texto extraído promedia menos que esto por página, se asume PDF escaneado/a mano y se cae
# al modo visión en vez de mandar un texto casi vacío.
UMBRAL_CARACTERES_POR_PAGINA = 40

SYSTEM_PROMPT_EXTRACCION = (
    "Eres un transcriptor de formularios. Se te entrega, en JSON, el ESQUEMA completo de un "
    "instrumento (secciones, preguntas con su id real, tipo, y — según el tipo — sus opciones, "
    "filas y columnas, cada una con su id real), y el contenido de un documento (PDF o Word) que "
    "alguien ya diligenció a mano o en computador fuera de este sistema. Tu única tarea es leer "
    "ese documento y transcribir, para cada campo que SÍ tenga contenido, una entrada que apunte "
    "al id real de la pregunta correspondiente en el esquema.\n\n"

    "=== REGLAS ESTRICTAS ===\n"
    "1. USA SOLO los ids que te dimos en el esquema — nunca inventes uno, nunca uses un id de "
    "otra pregunta/opción/fila/columna por parecido de texto.\n"
    "2. Si un campo del documento está en blanco, no se entiende, o no encuentras dónde quedó en "
    "el esquema, SIMPLEMENTE NO incluyas una entrada para él — no inventes contenido para "
    "completar, no dejes texto_libre vacío como relleno.\n"
    "3. Pregunta tipo 'abierta': una entrada con texto_libre = lo transcrito literalmente (puedes "
    "corregir errores obvios de tipeo, pero no resumas ni parafrasees el contenido real).\n"
    "4. Pregunta tipo 'unica'/'multiple': una entrada con opciones = [id de la(s) opción(es) "
    "marcada(s)], usando el id de la opción cuyo texto mejor coincida con lo marcado/escrito — "
    "nunca opciones no presentes en el esquema de esa pregunta.\n"
    "5. Pregunta tipo 'matriz': UNA entrada POR CADA celda (fila×columna) que tenga contenido, "
    "con fila y columna = los ids correspondientes y texto_libre = lo escrito en esa celda "
    "(si la celda es de semáforo 🟢🟡🔴 y la columna es justo esa, texto_libre puede ser el color "
    "o símbolo marcado). Una matriz típica del semáforo tiene una fila por dimensión y una "
    "columna por cada aspecto (color, cómo funciona hoy, problema/brecha, propuesta) — no mezcles "
    "el contenido de una columna en otra.\n"
    "6. Nunca agregues fila/columna a una pregunta que no sea tipo matriz, ni opciones a una "
    "pregunta abierta, ni texto_libre a una pregunta de opción (salvo que sea matriz).\n"
    "7. RESPONSABLE. Además de las respuestas, busca quién diligenció el documento — casi siempre "
    "está en las primeras líneas o en un encabezado, con etiquetas como 'Responsable', "
    "'Diligenciado por', 'Nombre', 'Elaborado por' o una firma al pie. Devuélvelo en el "
    "bloque 'responsable', copiando el texto TAL CUAL aparece (no lo normalices, no lo "
    "completes, no deduzcas un correo que no esté escrito). Si un dato no está, déjalo en null; "
    "si no encuentras ningún responsable, deja todo el bloque en null. NUNCA inventes un nombre "
    "ni uses el de una persona mencionada dentro de una respuesta — solo quien firma o declara "
    "haber diligenciado el documento.\n\n"

    "=== FORMATO DE SALIDA (obligatorio) ===\n"
    "Responde ÚNICAMENTE con un objeto JSON válido, sin explicación antes ni después, sin fences "
    "de markdown, con esta forma exacta:\n"
    "{\n"
    '  "responsable": {"nombre": "<texto o null>", "correo": "<texto o null>", '
    '"cargo": "<texto o null>", "dependencia": "<texto o null>"},\n'
    '  "respuestas": [\n'
    '    {"pregunta": <id>, "texto_libre": "<texto o \'\'>", "opciones": [<ids>], '
    '"fila": <id o null>, "columna": <id o null>}\n'
    "  ]\n"
    "}\n"
)


def _extraer_texto_docx(archivo):
    from docx import Document

    documento = Document(archivo)
    partes = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            if any(celdas):
                partes.append(' | '.join(celdas))
    return '\n'.join(partes), None


def _extraer_texto_o_imagenes_pdf(archivo):
    """Devuelve (texto, None) si el PDF tiene texto seleccionable suficiente, o (None,
    lista_de_imagenes_base64) si hay que caer a visión (escaneado/a mano)."""
    import pdfplumber

    archivo.seek(0)
    paginas_texto = []
    with pdfplumber.open(archivo) as pdf:
        n_paginas = len(pdf.pages)
        for pagina in pdf.pages:
            paginas_texto.append(pagina.extract_text() or '')

    texto = '\n'.join(paginas_texto).strip()
    if n_paginas and len(texto) / n_paginas >= UMBRAL_CARACTERES_POR_PAGINA:
        return texto, None

    import fitz  # PyMuPDF

    archivo.seek(0)
    imagenes = []
    documento = fitz.open(stream=archivo.read(), filetype='pdf')
    for pagina in documento[:MAX_PAGINAS_IMAGEN]:
        pixmap = pagina.get_pixmap(dpi=150)
        imagenes.append(base64.b64encode(pixmap.tobytes('png')).decode('ascii'))
    documento.close()
    return None, imagenes


def _leer_documento(extraccion):
    """Despacha por extensión de archivo. Devuelve (texto_o_None, imagenes_base64_o_None)."""
    nombre = extraccion.nombre_archivo_original or extraccion.archivo.name
    extension = nombre.rsplit('.', 1)[-1].lower() if '.' in nombre else ''

    extraccion.archivo.open('rb')
    try:
        if extension == 'docx':
            texto, _ = _extraer_texto_docx(extraccion.archivo)
            return texto, None
        if extension == 'pdf':
            return _extraer_texto_o_imagenes_pdf(extraccion.archivo)
        raise ValueError(f'Formato de archivo no soportado: .{extension} (solo .pdf o .docx).')
    finally:
        extraccion.archivo.close()


def _construir_payload_esquema(instrumento):
    secciones = []
    for seccion in instrumento.secciones.filter(activa=True).order_by('orden'):
        if seccion.tipo != seccion.TIPO_PREGUNTAS:
            continue
        preguntas = []
        for pregunta in seccion.preguntas.filter(activa=True).order_by('orden'):
            preguntas.append({
                'id': pregunta.id,
                'texto': pregunta.texto,
                'tipo': pregunta.tipo,
                'opciones': [{'id': o.id, 'texto': o.texto} for o in pregunta.opciones.all()],
                'filas': [{'id': f.id, 'texto': f.texto} for f in pregunta.filas.all()],
                'columnas': [{'id': c.id, 'texto': c.texto} for c in pregunta.columnas.all()],
            })
        secciones.append({'titulo': seccion.titulo, 'preguntas': preguntas})
    return {'instrumento': instrumento.nombre, 'secciones': secciones}


def _llamar_openai_extraccion(esquema, texto_documento=None, imagenes_base64=None):
    """Una sola llamada a OpenAI en modo JSON estricto, con texto o con imágenes (visión). Devuelve
    (dict_o_None, error) — nunca lanza excepción."""
    api_key = os.environ.get('OPENAI_API_KEY')
    if not api_key:
        return None, 'OPENAI_API_KEY no está configurada en el entorno del servidor (.env).'

    encabezado = 'ESQUEMA DEL INSTRUMENTO (JSON):\n' + json.dumps(esquema, ensure_ascii=False, indent=2)

    if imagenes_base64:
        contenido_usuario = [
            {'type': 'text', 'text': encabezado + '\n\nDOCUMENTO DILIGENCIADO (páginas adjuntas como imagen):'},
        ]
        for imagen_b64 in imagenes_base64:
            contenido_usuario.append({
                'type': 'image_url',
                'image_url': {'url': f'data:image/png;base64,{imagen_b64}', 'detail': 'high'},
            })
    else:
        contenido_usuario = (
            encabezado + '\n\nDOCUMENTO DILIGENCIADO (texto extraído):\n' + (texto_documento or '')
        )

    resultado = {}

    def _run():
        try:
            from openai import OpenAI
            client = OpenAI(api_key=api_key)
            kwargs = dict(
                model=DEFAULT_MODEL,
                messages=[
                    {'role': 'system', 'content': SYSTEM_PROMPT_EXTRACCION},
                    {'role': 'user', 'content': contenido_usuario},
                ],
                max_completion_tokens=MAX_OUTPUT_TOKENS,
                response_format={'type': 'json_object'},
            )
            if REASONING_EFFORT:
                kwargs['reasoning_effort'] = REASONING_EFFORT
            else:
                kwargs['temperature'] = 0.2
            respuesta = client.chat.completions.create(**kwargs)
            resultado['texto'] = respuesta.choices[0].message.content.strip()
        except Exception as exc:  # noqa: BLE001 — cualquier falla de la API cae a error legible
            resultado['error'] = str(exc)

    hilo = threading.Thread(target=_run, daemon=True)
    hilo.start()
    hilo.join(timeout=GENERATION_TIMEOUT_SECONDS)

    if hilo.is_alive():
        return None, f'Tiempo de espera agotado ({GENERATION_TIMEOUT_SECONDS}s) esperando a OpenAI.'
    if resultado.get('error'):
        return None, resultado['error']
    texto = resultado.get('texto')
    if not texto:
        return None, 'OpenAI no devolvió contenido.'
    try:
        return json.loads(texto), None
    except json.JSONDecodeError as exc:
        return None, f'OpenAI devolvió JSON inválido: {exc}'


def _guardar_respuestas(aplicacion, items_crudos, preguntas_validas):
    """Valida y guarda cada entrada devuelta por la IA — reutiliza la misma validación por tipo
    que ya usa el envío normal desde la web (ver views_participante._validar_entrada_instrumento),
    pero a diferencia de ese endpoint NO aborta todo si una entrada falla: la extracción es
    aproximada por naturaleza, así que una celda mal formada se omite (queda en
    preguntas_omitidas) y el resto se guarda igual — todo cae en revisión humana de todos modos."""
    from .models import RespuestaInstrumento
    from .serializers_participante import RespuestaInstrumentoEnvioItemSerializer
    from .views_participante import _validar_entrada_instrumento

    omitidas = []
    guardadas = 0
    for crudo in items_crudos:
        pregunta_id = crudo.get('pregunta')
        pregunta = preguntas_validas.get(pregunta_id)
        if pregunta is None:
            omitidas.append(pregunta_id)
            continue

        # El prompt le pide a la IA mandar SIEMPRE "fila"/"columna" (null si no aplica) para que
        # nunca olvide el campo en preguntas tipo matriz — pero PrimaryKeyRelatedField(required=
        # False) sin allow_null=True rechaza un null EXPLÍCITO (solo acepta la ausencia de la
        # clave). Sin este descarte, toda pregunta no-matriz fallaría la validación.
        crudo = {
            clave: valor for clave, valor in crudo.items()
            if not (clave in ('fila', 'columna') and valor is None)
        }

        entrada = RespuestaInstrumentoEnvioItemSerializer(data=crudo)
        if not entrada.is_valid():
            omitidas.append(pregunta_id)
            continue
        item = entrada.validated_data

        try:
            _validar_entrada_instrumento(pregunta, item)
        except Exception:  # noqa: BLE001 — ValidationError u otra inconsistencia de la IA
            omitidas.append(pregunta_id)
            continue

        respuesta, _ = RespuestaInstrumento.objects.update_or_create(
            aplicacion=aplicacion,
            pregunta=pregunta,
            fila=item.get('fila'),
            columna=item.get('columna'),
            defaults={'texto_libre': item.get('texto_libre', '')},
        )
        respuesta.opciones.set(item.get('opciones', []))
        guardadas += 1

    return guardadas, omitidas


def _limpiar_responsable(crudo):
    """Normaliza el bloque `responsable` que devolvió la IA a un dict de strings. La IA puede
    mandar null, omitirlo, o poner cualquier cosa en los campos — nada de eso puede tumbar una
    extracción que por lo demás salió bien."""
    if not isinstance(crudo, dict):
        return {}
    limpio = {}
    for clave in ('nombre', 'correo', 'cargo', 'dependencia'):
        valor = crudo.get(clave)
        if isinstance(valor, str) and valor.strip():
            limpio[clave] = valor.strip()
    return limpio


def emparejar_responsable_instrumento(instrumento, responsable):
    """Busca a qué usuario corresponde el responsable que leyó la IA.

    Se intenta primero contra los **preregistrados de este instrumento** (el conjunto chico y
    correcto: a esa gente ya se le asignó este instrumento) y solo si ahí no hay nada contra el
    resto de usuarios no-staff. Ese orden importa: buscar de una en todos los usuarios haría
    ambiguo cualquier nombre común, y un homónimo que no tiene nada que ver con el instrumento no
    debería competir con alguien a quien sí se le asignó."""
    from django.contrib.auth import get_user_model

    from .models import PreregistroInstrumento

    if not responsable:
        return None, emparejamiento.ESTADO_SIN_DATO

    nombre = responsable.get('nombre')
    correo = responsable.get('correo')

    preregistrados = [
        (pr.usuario, f'{pr.usuario.first_name} {pr.usuario.last_name}'.strip() or pr.usuario.username,
         pr.usuario.email)
        for pr in PreregistroInstrumento.objects.filter(instrumento=instrumento).select_related('usuario')
    ]
    usuario, estado = emparejamiento.emparejar(preregistrados, nombre=nombre, correo=correo)
    if usuario is not None or estado == emparejamiento.ESTADO_AMBIGUO:
        return usuario, estado

    Usuario = get_user_model()
    ids_preregistrados = {u.id for u, _n, _c in preregistrados}
    otros = [
        (u, f'{u.first_name} {u.last_name}'.strip() or u.username, u.email)
        for u in Usuario.objects.filter(is_staff=False).exclude(id__in=ids_preregistrados)
    ]
    return emparejamiento.emparejar(otros, nombre=nombre, correo=correo)


def _escribir_aplicacion(extraccion, resultado):
    """Escribe la transcripción como una AplicacionInstrumento del usuario de la extracción.
    Separada de procesar_extraccion_instrumento porque es exactamente lo que hay que volver a
    correr cuando el responsable se asigna después (ver asignar_responsable_instrumento) — el
    mismo trabajo, con la transcripción que ya estaba guardada."""
    from .models import AplicacionInstrumento, ExtraccionInstrumento, PreguntaInstrumento, PreregistroInstrumento

    preregistro, _ = PreregistroInstrumento.objects.get_or_create(
        instrumento=extraccion.instrumento, usuario=extraccion.usuario,
        defaults={'creado_por': extraccion.solicitado_por},
    )
    aplicacion, _ = AplicacionInstrumento.objects.get_or_create(preregistro=preregistro)

    preguntas_validas = {
        p.id: p for p in PreguntaInstrumento.objects.filter(
            seccion__instrumento=extraccion.instrumento, seccion__activa=True, activa=True,
        )
    }
    _guardadas, omitidas = _guardar_respuestas(
        aplicacion, resultado.get('respuestas') or [], preguntas_validas
    )

    aplicacion.generado_por_ia = True
    aplicacion.estado = AplicacionInstrumento.ESTADO_PENDIENTE
    aplicacion.enviado_en = timezone.now()
    aplicacion.revisado_por = None
    aplicacion.revisado_en = None
    aplicacion.save()

    extraccion.aplicacion = aplicacion
    extraccion.preguntas_omitidas = omitidas
    extraccion.estado = ExtraccionInstrumento.ESTADO_COMPLETO
    extraccion.error_mensaje = ''
    extraccion.modelo_usado = MODELO_USADO_LABEL
    extraccion.completado_en = timezone.now()
    extraccion.save(update_fields=[
        'aplicacion', 'preguntas_omitidas', 'estado', 'error_mensaje', 'modelo_usado',
        'completado_en', 'usuario', 'responsable_estado', 'resultado_crudo',
    ])
    return aplicacion


def asignar_responsable_instrumento(extraccion, usuario):
    """Termina una extracción que quedó en `sin_responsable`: le asigna la persona y escribe la
    AplicacionInstrumento con la transcripción que ya estaba guardada, sin volver a llamar a
    OpenAI. `responsable_estado` NO se toca — deja constancia de por qué hubo que asignar a mano
    (sin coincidencia, ambiguo, el documento no traía responsable)."""
    extraccion.usuario = usuario
    return _escribir_aplicacion(extraccion, extraccion.resultado_crudo or {})


def procesar_extraccion_instrumento(extraccion_id):
    """Genera la AplicacionInstrumento de una ExtraccionInstrumento ya creada (estado
    'pendiente'). Corre en un hilo de background — mismo patrón que
    analitica.analisis_ia_openai.analizar_momento_ia."""
    from django.db import close_old_connections

    close_old_connections()
    from .models import AplicacionInstrumento, ExtraccionInstrumento, PreguntaInstrumento, PreregistroInstrumento

    extraccion = None
    try:
        extraccion = ExtraccionInstrumento.objects.select_related('instrumento', 'usuario').get(pk=extraccion_id)
        extraccion.estado = ExtraccionInstrumento.ESTADO_PROCESANDO
        extraccion.save(update_fields=['estado'])

        texto, imagenes = _leer_documento(extraccion)
        esquema = _construir_payload_esquema(extraccion.instrumento)
        resultado, error = _llamar_openai_extraccion(esquema, texto_documento=texto, imagenes_base64=imagenes)

        if resultado is None:
            extraccion.estado = ExtraccionInstrumento.ESTADO_ERROR
            extraccion.error_mensaje = error or 'Error desconocido generando la extracción.'
            extraccion.save(update_fields=['estado', 'error_mensaje'])
            return

        # La transcripción se guarda SIEMPRE, antes de saber a quién atribuirla: es el trabajo
        # caro (una llamada a OpenAI con el documento entero) y no puede perderse porque el
        # responsable no se haya podido emparejar.
        extraccion.resultado_crudo = resultado

        if extraccion.usuario_id is None:
            detectado = _limpiar_responsable(resultado.get('responsable'))
            usuario, estado_responsable = emparejar_responsable_instrumento(
                extraccion.instrumento, detectado,
            )
            extraccion.responsable_detectado = detectado
            extraccion.responsable_estado = estado_responsable
            extraccion.usuario = usuario

        if extraccion.usuario_id is None:
            # Transcrito pero sin dueño: queda esperando que un admin asigne a la persona con
            # asignar_responsable_instrumento, que termina exactamente este mismo trabajo. No se
            # crea una cuenta a partir del nombre leído — ver jornadas/emparejamiento.py.
            extraccion.estado = ExtraccionInstrumento.ESTADO_SIN_RESPONSABLE
            extraccion.error_mensaje = ''
            extraccion.modelo_usado = MODELO_USADO_LABEL
            extraccion.save(update_fields=[
                'resultado_crudo', 'responsable_detectado', 'responsable_estado', 'usuario',
                'estado', 'error_mensaje', 'modelo_usado',
            ])
            return

        _escribir_aplicacion(extraccion, resultado)
    except Exception as exc:  # noqa: BLE001 — nunca debe dejar el hilo morir en silencio
        if extraccion is not None:
            extraccion.estado = ExtraccionInstrumento.ESTADO_ERROR
            extraccion.error_mensaje = str(exc)
            extraccion.save(update_fields=['estado', 'error_mensaje'])
    finally:
        close_old_connections()
