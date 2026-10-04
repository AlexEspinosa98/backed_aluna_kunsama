"""Tests de `jornadas.lectura_documentos` — la lectura del documento diligenciado fuera de la web
y la partición del esquema en lotes. Nada de esto toca la base de datos ni OpenAI: se construye un
.docx en memoria con python-docx y se mira el texto que saldría hacia el modelo."""
import io

from django.test import SimpleTestCase

from . import lectura_documentos


def _docx_en_memoria(construir):
    """Un .docx armado al vuelo por `construir(documento)`, devuelto como buffer listo para leer."""
    from docx import Document

    documento = Document()
    construir(documento)
    buffer = io.BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    return buffer


def _tabla(documento, filas):
    tabla = documento.add_table(rows=len(filas), cols=len(filas[0]))
    for i, fila in enumerate(filas):
        for j, celda in enumerate(fila):
            tabla.cell(i, j).text = celda
    return tabla


class ExtraerTextoDocxTests(SimpleTestCase):
    """El bug que motivó este módulo: el lector anterior recorría todos los párrafos y DESPUÉS
    todas las tablas, así que el modelo recibía los títulos de sección separados de sus tablas."""

    def test_respeta_el_orden_del_documento(self):
        def construir(documento):
            documento.add_paragraph('Sección A')
            _tabla(documento, [['Variable', 'Valor'], ['Créditos', '3']])
            documento.add_paragraph('Sección B')
            _tabla(documento, [['Variable', 'Valor'], ['Créditos', '4']])

        texto = lectura_documentos.extraer_texto_docx(_docx_en_memoria(construir))
        lineas = [l for l in texto.split('\n') if l]

        self.assertLess(lineas.index('Sección A'), lineas.index('Sección B'))
        # La tabla de A va ANTES del título de B — con el lector anterior las dos tablas caían
        # juntas al final, después de los dos títulos.
        self.assertLess(
            next(i for i, l in enumerate(lineas) if 'Créditos | 3' in l),
            lineas.index('Sección B'),
        )

    def test_cada_tabla_queda_rotulada_con_el_parrafo_que_la_precede(self):
        """Lo que hace transcribibles las tres tablas idénticas del formato de diagnóstico de
        articulación académica: solo las distingue el párrafo "Componente N" de arriba."""
        filas = [['Variable', 'Programa A'], ['Nombre', '']]

        def construir(documento):
            documento.add_paragraph('Componente 1')
            _tabla(documento, filas)
            documento.add_paragraph('Componente 2')
            _tabla(documento, filas)

        texto = lectura_documentos.extraer_texto_docx(_docx_en_memoria(construir))

        self.assertIn('[TABLA 1 — justo debajo de "Componente 1"]', texto)
        self.assertIn('[TABLA 2 — justo debajo de "Componente 2"]', texto)
        self.assertIn('[fin TABLA 1]', texto)

    def test_las_filas_totalmente_vacias_no_se_mandan(self):
        def construir(documento):
            _tabla(documento, [['Profesor', 'Área'], ['', ''], ['Ana', 'Física']])

        texto = lectura_documentos.extraer_texto_docx(_docx_en_memoria(construir))

        self.assertIn('Ana | Física', texto)
        self.assertNotIn('\n | \n', texto)

    def test_un_docx_sin_nada_devuelve_texto_vacio(self):
        """La señal que usan los dos extractores para cortar antes de gastar una llamada al
        proveedor (ver `procesar_extraccion_instrumento`)."""
        texto = lectura_documentos.extraer_texto_docx(_docx_en_memoria(lambda documento: None))
        self.assertEqual(texto.strip(), '')

    def test_rechaza_una_extension_no_soportada(self):
        with self.assertRaises(ValueError):
            lectura_documentos.leer_documento(io.BytesIO(b'x'), 'formato.xlsx')


class AgruparEnLotesTests(SimpleTestCase):
    """La partición que evita que la respuesta del modelo se trunque."""

    def test_agrupa_hasta_llenar_el_presupuesto_sin_reordenar(self):
        lotes = lectura_documentos.agrupar_en_lotes(
            ['a', 'b', 'c', 'd'], costo=lambda _item: 3, presupuesto=7,
        )
        self.assertEqual(lotes, [['a', 'b'], ['c', 'd']])

    def test_un_item_mas_grande_que_el_presupuesto_va_solo(self):
        """No se parte por dentro: una matriz de 14×4 es una sola pregunta y el prompt necesita
        verla completa con todas sus filas y columnas."""
        lotes = lectura_documentos.agrupar_en_lotes(
            ['chico', 'enorme', 'chico2'],
            costo=lambda item: 100 if item == 'enorme' else 1,
            presupuesto=10,
        )
        self.assertEqual(lotes, [['chico'], ['enorme'], ['chico2']])

    def test_sin_items_no_hay_lotes(self):
        self.assertEqual(lectura_documentos.agrupar_en_lotes([], lambda _i: 1, 10), [])


class TranscribirEnPartesTests(SimpleTestCase):
    """La red ante un truncado: el lote se parte en dos y se reintenta, en vez de morir y obligar a
    volver a subir el documento. Salió de un fallo real en producción (2026-10-03): un .docx de
    47.500 caracteres con respuestas largas truncó el bloque 4 de 5."""

    def test_sin_truncado_es_una_sola_llamada(self):
        llamadas = []

        def llamar(lote):
            llamadas.append(tuple(lote))
            return {'respuestas': list(lote)}, None, False

        resultados, error = lectura_documentos.transcribir_en_partes(
            [1, 2, 3, 4], llamar, describir=str,
        )
        self.assertIsNone(error)
        self.assertEqual(llamadas, [(1, 2, 3, 4)])
        self.assertEqual(resultados, [{'respuestas': [1, 2, 3, 4]}])

    def test_un_truncado_parte_el_lote_en_dos_y_reintenta(self):
        llamadas = []

        def llamar(lote):
            llamadas.append(tuple(lote))
            if len(lote) == 4:
                return None, 'se truncó', True
            return {'respuestas': list(lote)}, None, False

        resultados, error = lectura_documentos.transcribir_en_partes(
            [1, 2, 3, 4], llamar, describir=str,
        )
        self.assertIsNone(error)
        self.assertEqual(llamadas, [(1, 2, 3, 4), (1, 2), (3, 4)])
        # Las dos mitades vuelven en orden, así el `responsable` del encabezado se fusiona
        # quedándose con el de la primera.
        self.assertEqual(resultados, [{'respuestas': [1, 2]}, {'respuestas': [3, 4]}])

    def test_parte_cuantas_veces_haga_falta(self):
        def llamar(lote):
            if len(lote) > 1:
                return None, 'se truncó', True
            return {'respuestas': list(lote)}, None, False

        resultados, error = lectura_documentos.transcribir_en_partes(
            [1, 2, 3, 4, 5], llamar, describir=str,
        )
        self.assertIsNone(error)
        self.assertEqual([r['respuestas'][0] for r in resultados], [1, 2, 3, 4, 5])

    def test_una_pregunta_sola_que_no_cabe_falla_nombrandola(self):
        """Ya no hay nada que partir: el error tiene que decir cuál es la pregunta, porque el
        arreglo pasa por acortarla o por subir el tope de tokens, no por reintentar."""
        def llamar(_lote):
            return None, 'se truncó', True

        resultados, error = lectura_documentos.transcribir_en_partes(
            ['matriz gigante'], llamar, describir=lambda p: f'la pregunta «{p}»',
        )
        self.assertIsNone(resultados)
        self.assertIn('No se puede partir más', error)
        self.assertIn('matriz gigante', error)

    def test_un_error_que_no_es_truncado_no_se_reintenta(self):
        llamadas = []

        def llamar(lote):
            llamadas.append(tuple(lote))
            return None, 'OPENAI_API_KEY no está configurada', False

        resultados, error = lectura_documentos.transcribir_en_partes(
            [1, 2, 3, 4], llamar, describir=str,
        )
        self.assertIsNone(resultados)
        self.assertEqual(error, 'OPENAI_API_KEY no está configurada')
        self.assertEqual(llamadas, [(1, 2, 3, 4)], 'no debe reintentar un fallo que no es truncado')


def _pdf_en_memoria(bloques):
    """Un PDF armado al vuelo con reportlab (ya es dependencia del proyecto). `bloques` es una
    lista de str (párrafo) o de list-de-filas (tabla con grilla, que es lo que pdfplumber necesita
    para detectarla)."""
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER)
    estilo = getSampleStyleSheet()['Normal']
    flujo = []
    for bloque in bloques:
        if isinstance(bloque, str):
            flujo.append(Paragraph(bloque, estilo))
        else:
            tabla = Table(bloque)
            tabla.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, (0, 0, 0))]))
            flujo.append(tabla)
    doc.build(flujo)
    buffer.seek(0)
    return buffer


class ExtraerTextoPdfTests(SimpleTestCase):
    """El lector de PDF reconstruyendo tablas. Antes usaba `extract_text()` a secas, que las
    aplana: la etiqueta de la fila y el valor de la columna quedaban pegados en un mismo renglón,
    sin encabezado ni separador. Medido sobre dos documentos reales, el acuerdo entre dos modelos
    distintos era 100% en .docx y 85% en PDF — la brecha la causaba este lector."""

    def test_reconstruye_la_tabla_con_separadores(self):
        filas = [['Variable', 'Programa A'], ['Créditos', '4'], ['Nombre', 'Cálculo I']]
        texto, imagenes = lectura_documentos.extraer_texto_o_imagenes_pdf(
            _pdf_en_memoria(['Un párrafo de contexto suficientemente largo para no caer a visión.',
                             filas])
        )
        self.assertIsNone(imagenes, 'con capa de texto no debe caer al modo visión')
        self.assertIn('Créditos | 4', texto)
        self.assertIn('Nombre | Cálculo I', texto)

    def test_cada_tabla_queda_rotulada_con_el_texto_que_la_precede(self):
        filas = [['Variable', 'Programa A'], ['Nombre', '']]
        texto, _ = lectura_documentos.extraer_texto_o_imagenes_pdf(_pdf_en_memoria([
            'Texto introductorio del formato, con largo suficiente para el umbral de visión.',
            'Componente 1', filas, 'Componente 2', filas,
        ]))
        self.assertIn('[TABLA 1 — justo debajo de "Componente 1"]', texto)
        self.assertIn('[TABLA 2 — justo debajo de "Componente 2"]', texto)
        self.assertIn('[fin TABLA 1]', texto)

    def test_el_texto_de_la_tabla_no_se_repite_como_texto_suelto(self):
        """Las líneas que caen dentro del recuadro de una tabla salen SOLO como filas de la tabla:
        si además salieran como texto suelto, el modelo vería el mismo dato dos veces."""
        filas = [['Dimensión', 'Hallazgo'], ['Saber Pro', 'Coordinado con Vicerrectoría']]
        texto, _ = lectura_documentos.extraer_texto_o_imagenes_pdf(_pdf_en_memoria([
            'Encabezado del instrumento con texto suficiente para pasar el umbral de visión.',
            filas,
        ]))
        self.assertEqual(texto.count('Coordinado con Vicerrectoría'), 1)


class TablasAnidadasYCorridasTests(SimpleTestCase):
    """Los dos filtros que hicieron falta al probar contra un PDF real de producción."""

    class _Tabla:
        def __init__(self, bbox):
            self.bbox = bbox

    def test_descarta_la_tabla_anidada_y_conserva_la_externa(self):
        externa = self._Tabla((50, 100, 500, 400))
        interna = self._Tabla((200, 150, 400, 300))
        quedan = lectura_documentos._sin_tablas_anidadas([externa, interna])
        self.assertEqual(quedan, [externa])

    def test_dos_tablas_que_no_se_solapan_se_conservan_las_dos(self):
        a = self._Tabla((50, 100, 500, 200))
        b = self._Tabla((50, 250, 500, 400))
        self.assertEqual(len(lectura_documentos._sin_tablas_anidadas([a, b])), 2)

    def test_dos_tablas_con_el_mismo_recuadro_dejan_una(self):
        a = self._Tabla((50, 100, 500, 200))
        b = self._Tabla((50, 100, 500, 200))
        self.assertEqual(len(lectura_documentos._sin_tablas_anidadas([a, b])), 1)


class UnirFilasTests(SimpleTestCase):
    """Los tres casos de una tabla partida por un salto de página, tal como aparecen en un PDF
    real del instrumento de diagnóstico."""

    def test_descarta_el_encabezado_repetido(self):
        """Word repite la fila de título cuando la tabla tiene esa opción; si no se descarta, queda
        como una fila de datos cuyos valores son los nombres de las columnas."""
        acumuladas = [['Dimensión', 'Hallazgo'], ['Saber Pro', 'Coordinado']]
        nuevas = [['Dimensión', 'Hallazgo'], ['Nivelación', 'Sin instancia']]
        self.assertEqual(
            lectura_documentos._unir_filas(acumuladas, nuevas),
            [['Dimensión', 'Hallazgo'], ['Saber Pro', 'Coordinado'], ['Nivelación', 'Sin instancia']],
        )

    def test_pega_la_fila_cortada_a_la_anterior(self):
        """Una fila cuyo contenido siguió en la página siguiente llega con la primera celda vacía.
        Abrirle una fila nueva es lo que hacía que el modelo transcribiera un profesor fantasma."""
        acumuladas = [['Profesor(a)', 'Formación'], ['Borish Cuadrado', 'Biólogo; MSc en']]
        nuevas = [['', 'Ciencias Ambientales'], ['Cindi Güete', 'Bióloga']]
        self.assertEqual(
            lectura_documentos._unir_filas(acumuladas, nuevas),
            [['Profesor(a)', 'Formación'],
             ['Borish Cuadrado', 'Biólogo; MSc en Ciencias Ambientales'],
             ['Cindi Güete', 'Bióloga']],
        )

    def test_filas_nuevas_se_agregan_tal_cual(self):
        acumuladas = [['Proceso', 'Quién'], ['Microdiseños', 'Jefe']]
        nuevas = [['Evaluación', 'Comité']]
        self.assertEqual(
            lectura_documentos._unir_filas(acumuladas, nuevas),
            [['Proceso', 'Quién'], ['Microdiseños', 'Jefe'], ['Evaluación', 'Comité']],
        )

    def test_no_pisa_lo_ya_escrito_al_pegar_una_fila_cortada(self):
        acumuladas = [['A', 'B'], ['uno', 'dos']]
        nuevas = [['', 'tres']]
        self.assertEqual(
            lectura_documentos._unir_filas(acumuladas, nuevas),
            [['A', 'B'], ['uno', 'dos tres']],
        )


class EsContinuacionTests(SimpleTestCase):
    """El criterio para decidir que una tabla continúa a la anterior. Los umbrales se calibraron
    contra las 21 tablas de un PDF real: tienen que unir las continuaciones verdaderas y NO unir
    dos tablas distintas que simplemente se suceden."""

    class _Tabla:
        def __init__(self, bbox, cols=(47, 150, 300, 450)):
            self.bbox = bbox
            self.cells = [(x, bbox[1], x + 50, bbox[3]) for x in cols]

    def _bloque(self, pagina, top, bottom, cols=(47, 150, 300, 450), alto=792):
        return {'pagina': pagina, 'top': top, 'alto': alto,
                'tabla': self._Tabla((47, top, 565, bottom), cols)}

    def test_une_la_que_llega_al_fondo_con_la_que_arranca_arriba(self):
        anterior = self._bloque(1, 353, 742)     # 94% de la página
        siguiente = self._bloque(2, 40, 740)     # 5% de la página
        self.assertTrue(lectura_documentos._es_continuacion(anterior, siguiente))

    def test_no_une_si_la_anterior_termina_a_media_pagina(self):
        """Una tabla que TERMINA a media página no se quedó sin espacio: lo que viene después es
        otra tabla. Es el caso de Componente 1 → Componente 2 en el formato real."""
        anterior = self._bloque(10, 40, 492)     # 62%, le sobraba página
        siguiente = self._bloque(11, 55, 738)
        self.assertFalse(lectura_documentos._es_continuacion(anterior, siguiente))

    def test_no_une_tablas_de_la_misma_pagina(self):
        anterior = self._bloque(7, 415, 737)
        siguiente = self._bloque(7, 40, 372)
        self.assertFalse(lectura_documentos._es_continuacion(anterior, siguiente))

    def test_no_une_si_las_columnas_no_coinciden(self):
        anterior = self._bloque(1, 353, 742, cols=(47, 150, 300, 450))
        siguiente = self._bloque(2, 40, 740, cols=(47, 200, 380))
        self.assertFalse(lectura_documentos._es_continuacion(anterior, siguiente))

    def test_no_une_paginas_no_consecutivas(self):
        self.assertFalse(lectura_documentos._es_continuacion(
            self._bloque(1, 353, 742), self._bloque(3, 40, 740)))


class PdfTablaQueCruzaPaginaTests(SimpleTestCase):
    """Integración: un PDF de verdad con una tabla que no cabe en una página."""

    def test_sale_como_una_sola_tabla_con_su_encabezado(self):
        filas = [['Profesor(a)', 'Formación']]
        filas += [[f'Docente {i}', f'Magíster número {i}'] for i in range(1, 61)]
        texto, imagenes = lectura_documentos.extraer_texto_o_imagenes_pdf(
            _pdf_en_memoria(['Encabezado del instrumento, con texto suficiente para el umbral.', filas])
        )
        self.assertIsNone(imagenes)
        self.assertEqual(texto.count('[TABLA '), 1, 'la tabla partida debe salir unificada')
        # Las filas de las dos páginas, bajo el mismo encabezado.
        self.assertIn('Profesor(a) | Formación', texto)
        self.assertIn('Docente 1 | Magíster número 1', texto)
        self.assertIn('Docente 60 | Magíster número 60', texto)
