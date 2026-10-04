"""Lectura del documento que alguien diligenció FUERA de la web (.docx o .pdf) y partición del
esquema en lotes, para las dos extracciones con IA del proyecto:
`instrumentos.extraccion_ia_openai` (módulo `instrumentos`) y
`participantes.extraccion_momento_ia_openai` (modelo clásico `jornadas.Momento`/`Pregunta`).

Vive en `jornadas` por lo mismo que `jornadas.emparejamiento`: es la app que las dos importan sin
armar un ciclo.

Este código estaba duplicado —idéntico— en los dos módulos, y esa duplicación costó un bug real:
el lector de .docx recorría `documento.paragraphs` primero y `documento.tables` después, así que
el texto que llegaba al modelo traía TODOS los títulos de sección al principio y TODAS las tablas
juntas al final, sin ninguna forma de saber cuál tabla iba debajo de cuál título. En el formato de
diagnóstico de articulación académica (UNIMAGDALENA) eso es fatal: sus tres tablas del análisis de
coherencia son IDÉNTICAS —mismos encabezados, mismas filas de variables— y lo único que las
distingue es el párrafo que las precede ("Componente 1", "Componente 2", "Componente 3"). Con el
orden perdido, transcribir esas 84 celdas al componente correcto era adivinar.

Ahora se recorre el cuerpo del documento en su orden real y cada tabla sale delimitada y rotulada
con el párrafo que la antecede."""

# Tope de páginas enviadas como imagen — un PDF escaneado de 90 preguntas rara vez pasa de esto, y
# limita el costo/tiempo de una sola llamada con visión.
MAX_PAGINAS_IMAGEN = 20
# Si el texto extraído de un PDF promedia menos que esto por página, se asume escaneado/a mano y se
# cae al modo visión en vez de mandar un texto casi vacío.
UMBRAL_CARACTERES_POR_PAGINA = 40
# Recorte del párrafo que rotula una tabla: alcanza para "Componente 2" o "3. Semáforo de
# articulación académica" sin arrastrar un párrafo de instrucciones completo dentro del rótulo.
MAX_CARACTERES_ROTULO = 120


def _bloques_en_orden(documento):
    """Recorre el cuerpo del .docx en su orden real, intercalando párrafos y tablas.

    `documento.paragraphs` y `documento.tables` son dos listas separadas: cada una respeta el orden
    entre sus propios elementos, pero no hay forma de saber cómo se intercalan entre sí. La única
    fuente de ese orden es el XML del cuerpo, donde `w:p` y `w:tbl` son hermanos."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for hijo in documento.element.body.iterchildren():
        if isinstance(hijo, CT_P):
            yield Paragraph(hijo, documento)
        elif isinstance(hijo, CT_Tbl):
            yield Table(hijo, documento)


def _filas_con_contenido(filas_de_celdas):
    """Las filas con algún contenido, como `celda | celda | celda`.

    Las filas completamente vacías se omiten: en un formato en blanco son cientos de renglones de
    `| | | |` que no le dicen nada al modelo y sí le gastan contexto. Una fila realmente
    diligenciada tiene, por definición, al menos una celda con texto."""
    for celdas in filas_de_celdas:
        limpias = [' '.join((c or '').split()) for c in celdas]
        if any(limpias):
            yield ' | '.join(limpias)


def _bloque_tabla(numero, rotulo, filas):
    """El rótulo y los delimitadores de una tabla. Compartido por los lectores de .docx y de PDF a
    propósito: los dos tienen que producir EXACTAMENTE el mismo formato, porque el prompt de
    extracción es uno solo y no sabe de qué tipo de archivo vino el texto."""
    encabezado = f'[TABLA {numero}'
    if rotulo:
        encabezado += f' — justo debajo de "{rotulo}"'
    return [encabezado + ']', *filas, f'[fin TABLA {numero}]']


def _rotulo(texto):
    return ' '.join(texto.split())[:MAX_CARACTERES_ROTULO]


def extraer_texto_docx(archivo):
    """Texto plano del .docx respetando el orden del documento, con cada tabla delimitada y
    rotulada con el párrafo que la precede.

    El rótulo es la pieza que le permite al modelo saber a qué pregunta del esquema corresponde una
    tabla cuando hay varias con encabezados idénticos (ver el docstring del módulo). Los
    delimitadores `[TABLA n]` / `[fin TABLA n]` cierran el otro agujero del formato plano anterior:
    sin ellos, dos tablas consecutivas con las mismas columnas se leen como una sola tabla larga."""
    from docx import Document

    documento = Document(archivo)
    partes = []
    rotulo_pendiente = ''
    numero_tabla = 0

    for bloque in _bloques_en_orden(documento):
        if hasattr(bloque, 'rows'):  # docx.table.Table
            numero_tabla += 1
            filas = list(_filas_con_contenido(
                [c.text for c in fila.cells] for fila in bloque.rows
            ))
            if filas:
                partes.extend(_bloque_tabla(numero_tabla, rotulo_pendiente, filas))
            continue

        texto = bloque.text.strip()
        if not texto:
            continue
        partes.append(texto)
        # Solo el último párrafo no vacío antes de la tabla la rotula: en este tipo de formato ese
        # párrafo es siempre el título de la sección o el del bloque ("Componente 2").
        rotulo_pendiente = _rotulo(texto)

    return '\n'.join(partes)


def _contenida(interna, externa, tolerancia=2):
    xi0, ti0, xi1, ti1 = interna
    xe0, te0, xe1, te1 = externa
    return (xi0 >= xe0 - tolerancia and ti0 >= te0 - tolerancia
            and xi1 <= xe1 + tolerancia and ti1 <= te1 + tolerancia)


def _sin_tablas_anidadas(tablas):
    """Descarta las tablas cuyo recuadro cae DENTRO del de otra, quedándose con la de afuera.

    `find_tables()` detecta a la vez la tabla externa y la que alguien anidó dentro de una de sus
    celdas (pasa en este formato: la columna "Hallazgo / evidencia" trae su propia tablita). Sin
    este filtro el mismo contenido llega dos veces, en dos tablas numeradas distinto, y el modelo
    lo transcribe dos veces o elige la equivocada. Se conserva la externa porque es la única que
    trae las etiquetas de fila; lo anidado sigue estando, aplanado dentro de su celda."""
    descartadas = set()
    for i, interna in enumerate(tablas):
        for j, externa in enumerate(tablas):
            if i != j and j not in descartadas and _contenida(interna.bbox, externa.bbox):
                descartadas.add(i)
                break
    return [t for i, t in enumerate(tablas) if i not in descartadas]


def _lineas_repetidas_en_cada_pagina(pdf, umbral=0.6):
    """Las líneas que aparecen en la mayoría de las páginas: encabezados y pies corridos.

    Importan por el rótulo de las tablas. En un PDF, la última línea antes de una tabla casi
    siempre es el encabezado de página ("Instrumento de trabajo · Articulación Académica · …"),
    no el título de la sección — así que sin filtrarlas TODAS las tablas quedan rotuladas igual y
    el rótulo no distingue nada, que es justo lo que tiene que hacer."""
    from collections import Counter

    paginas = len(pdf.pages)
    if paginas < 4:
        return set()
    conteo = Counter()
    for pagina in pdf.pages:
        try:
            lineas = pagina.extract_text_lines()
        except Exception:  # noqa: BLE001
            continue
        conteo.update({' '.join((l['text'] or '').split()) for l in lineas})
    return {texto for texto, veces in conteo.items()
            if texto and veces >= max(3, int(paginas * umbral))}


def _dentro_de_alguna_tabla(linea, bboxes):
    """Si el centro vertical de la línea cae dentro del recuadro de una tabla (y se solapa
    horizontalmente con ella), la línea es contenido de esa tabla y no texto suelto."""
    centro = (linea['top'] + linea['bottom']) / 2
    for x0, top, x1, bottom in bboxes:
        if top <= centro <= bottom and linea['x1'] > x0 and linea['x0'] < x1:
            return True
    return False


def _texto_pdf_estructurado(pdf):
    """Texto del PDF en orden de lectura, con las tablas RECONSTRUIDAS, delimitadas y rotuladas —
    el mismo formato que produce `extraer_texto_docx`.

    `pagina.extract_text()` a secas (lo único que se usaba antes) aplana las tablas: las celdas
    salen como líneas de texto sueltas, el nombre de la fila y el valor de la columna quedan
    pegados en un mismo renglón, y no hay encabezado ni separador. En un formato de diagnóstico
    real eso llegaba así al modelo:

        Licenciatura en Química Biología Celular
        Biología              Biología General
        Química               Biología Celular

    En una tabla de dos columnas el modelo puede adivinar; en el semáforo de cinco no tiene cómo.
    Medido sobre dos documentos reales con los mismos dos modelos, el acuerdo de transcripción era
    del 100% en .docx y del 85% en PDF, y las discrepancias eran justamente celdas con tres
    valores de columnas distintas pegados — la brecha la causaba este lector, no el modelo.

    Acá se usa `find_tables()` para ubicar los recuadros, `extract()` para sacar las celdas, y
    `extract_text_lines()` para el texto de fuera; todo se ordena por posición vertical para
    respetar el orden de lectura. Los saltos de línea DENTRO de una celda se colapsan a un espacio:
    en un PDF son artefactos de maquetación (la celda era angosta), no parte de la respuesta."""
    partes = []
    rotulo_pendiente = ''
    numero_tabla = 0
    corridas = _lineas_repetidas_en_cada_pagina(pdf)

    for pagina in pdf.pages:
        tablas = _sin_tablas_anidadas(pagina.find_tables())
        # Los recuadros para excluir texto son los de TODAS las tablas detectadas, no solo las que
        # se conservan: el contenido de una tabla anidada sigue siendo contenido de tabla y no
        # debe volver a salir como texto suelto.
        bboxes = [t.bbox for t in pagina.find_tables()]

        bloques = [(t.bbox[1], 'tabla', t) for t in tablas]
        try:
            lineas = pagina.extract_text_lines()
        except Exception:  # noqa: BLE001 — pdfplumber viejo o página rara: se cae a texto plano
            lineas = []
            texto_plano = (pagina.extract_text() or '').strip()
            if texto_plano:
                partes.append(texto_plano)
        for linea in lineas:
            if _dentro_de_alguna_tabla(linea, bboxes):
                continue
            if ' '.join((linea['text'] or '').split()) in corridas:
                continue
            bloques.append((linea['top'], 'texto', linea['text']))

        for _tope, tipo, dato in sorted(bloques, key=lambda b: b[0]):
            if tipo == 'tabla':
                numero_tabla += 1
                filas = list(_filas_con_contenido(dato.extract()))
                if filas:
                    partes.extend(_bloque_tabla(numero_tabla, rotulo_pendiente, filas))
                continue
            texto = (dato or '').strip()
            if not texto:
                continue
            partes.append(texto)
            rotulo_pendiente = _rotulo(texto)

    return '\n'.join(partes)


def extraer_texto_o_imagenes_pdf(archivo):
    """Devuelve `(texto, None)` si el PDF tiene texto seleccionable suficiente, o
    `(None, lista_de_imagenes_base64)` si hay que caer a visión (escaneado / diligenciado a mano)."""
    import base64

    import pdfplumber

    archivo.seek(0)
    with pdfplumber.open(archivo) as pdf:
        n_paginas = len(pdf.pages)
        # La decisión texto-vs-visión se toma con el texto plano, que es barato: si el PDF está
        # escaneado no hay tablas que reconstruir y el trabajo estructurado sería en vano.
        texto_plano = '\n'.join((p.extract_text() or '') for p in pdf.pages).strip()
        if n_paginas and len(texto_plano) / n_paginas >= UMBRAL_CARACTERES_POR_PAGINA:
            return _texto_pdf_estructurado(pdf), None

    import fitz  # PyMuPDF

    archivo.seek(0)
    imagenes = []
    documento = fitz.open(stream=archivo.read(), filetype='pdf')
    for pagina in documento[:MAX_PAGINAS_IMAGEN]:
        pixmap = pagina.get_pixmap(dpi=150)
        imagenes.append(base64.b64encode(pixmap.tobytes('png')).decode('ascii'))
    documento.close()
    return None, imagenes


def leer_documento(archivo, nombre):
    """Despacha por extensión del nombre original. Devuelve `(texto_o_None, imagenes_o_None)`.

    El archivo lo abre y cierra quien llama: acá no se sabe si es un `FieldFile` de Django o un
    buffer en memoria."""
    extension = nombre.rsplit('.', 1)[-1].lower() if '.' in nombre else ''
    if extension == 'docx':
        return extraer_texto_docx(archivo), None
    if extension == 'pdf':
        return extraer_texto_o_imagenes_pdf(archivo)
    raise ValueError(f'Formato de archivo no soportado: .{extension} (solo .pdf o .docx).')


def agrupar_en_lotes(items, costo, presupuesto):
    """Parte `items` en lotes CONSECUTIVOS (sin reordenar) cuyo `costo(item)` sumado no pase de
    `presupuesto`. Un item que por sí solo lo excede se va en un lote propio — partirlo por dentro
    no le toca a esta función.

    Es lo que evita que una transcripción se trunque: el tope de tokens de salida de una llamada no
    depende del tamaño del documento sino de cuántas celdas se le piden transcribir en ella, así
    que el control tiene que estar acá y no en el prompt."""
    lotes = []
    actual = []
    acumulado = 0
    for item in items:
        costo_item = max(1, costo(item))
        if actual and acumulado + costo_item > presupuesto:
            lotes.append(actual)
            actual = []
            acumulado = 0
        actual.append(item)
        acumulado += costo_item
    if actual:
        lotes.append(actual)
    return lotes


def transcribir_en_partes(lote, llamar, describir):
    """Llama `llamar(lote)` y, si la respuesta salió TRUNCADA, parte el lote en dos y reintenta
    cada mitad — recursivamente, hasta que cada parte quepa o quede una sola pregunta.

    `llamar(sublote)` devuelve `(resultado, error, truncado)`. Devuelve
    `(lista_de_resultados, None)` o `(None, error)`.

    Por qué esto y no solo un presupuesto más chico: el tamaño que cabe en una respuesta no se
    puede calcular de antemano. Depende de cuánto texto escribió quien diligenció el documento
    (un formato de 47.500 caracteres con respuestas largas gasta por celda muchísimo más que uno
    de respuestas de tres palabras) y, con un modelo de razonamiento, también de cuántos tokens
    gastó razonando, que salen del mismo tope. Cualquier constante que se elija va a ser
    demasiado grande para algún documento real. Así que el presupuesto queda como la apuesta
    inicial —la que evita reintentos en el caso normal— y esto es la red: ante un truncado se
    divide y se reintenta, sin que nadie tenga que volver a subir el archivo.

    `describir(item)` solo se usa para el mensaje de error del caso en que una ÚNICA pregunta no
    cabe: ahí ya no hay nada que partir y hay que decir cuál es."""
    resultado, error, truncado = llamar(lote)
    if resultado is not None:
        return [resultado], None
    if not truncado:
        return None, error
    if len(lote) <= 1:
        detalle = f' No se puede partir más: {describir(lote[0])} sola no cabe.' if lote else ''
        return None, f'{error}{detalle}'

    mitad = len(lote) // 2
    partes = []
    for sublote in (lote[:mitad], lote[mitad:]):
        resultados, error_parte = transcribir_en_partes(sublote, llamar, describir)
        if resultados is None:
            return None, error_parte
        partes.extend(resultados)
    return partes, None
